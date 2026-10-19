"""PDF -> DOCX PPTX XLSX JPG PNG EPUB TXT MD HTML (all local, PyMuPDF based)."""
from __future__ import annotations

import base64
import html as htmlmod
import io
import re
from collections import Counter
from pathlib import Path

from .. import formats
from ..i18n import tr
from ..paths import safe_filename
from ..registry import Option, UserError, add
from .documents import media_dir_name, pandoc, wrap_html

_IMG_MIN = 16  # ignore images smaller than this (px) when extracting


def open_pdf(src: Path):
    """Open a PDF for reading; UserError if it cannot be opened or is locked."""
    import pymupdf

    try:
        doc = pymupdf.open(str(src))
    except Exception as e:
        raise UserError(tr("Cannot open {name}: {error}").format(name=Path(src).name, error=e)) from e
    if doc.needs_pass:
        doc.close()
        raise UserError(tr("{name} is password protected — unlock it first").format(name=Path(src).name))
    if doc.page_count == 0:
        doc.close()
        raise UserError(tr("{name} has no pages").format(name=Path(src).name))
    return doc


def find_tables(page) -> list:
    """Tables PyMuPDF detects on `page` ([] if detection fails)."""
    import pymupdf

    if hasattr(pymupdf, "no_recommend_layout"):
        pymupdf.no_recommend_layout()  # it would otherwise print() a tip to stdout
    try:
        return list(page.find_tables().tables)
    except Exception:
        return []


# --------------------------------------------------------------------------- images


def render_pages(src: Path, ext: str, ctx, dpi: int | None = None, src_for_name: Path | None = None) -> Path:
    """Render every page; one file for a single page, else a folder '<name> pages'
    holding '<name> page 01.png' ... (both names in the UI language)."""
    from ..config import settings

    dpi = int(dpi or settings.get("image_dpi", 300))
    quality = int(settings.get("jpeg_quality", 90))
    name_src = Path(src_for_name or src)
    stem = Path(name_src.name).stem if name_src.is_dir() else name_src.stem
    with open_pdf(src) as doc:
        n = doc.page_count
        if n == 1:
            out = ctx.output_path(name_src, ext, name=stem)
            _save_pix(doc[0].get_pixmap(dpi=dpi), out, ext, quality)
            return out
        folder = ctx.output_path(name_src, "", name=safe_filename(tr("{name} pages").format(name=stem)))
        folder.mkdir(parents=True)
        pad = max(2, len(str(n)))
        for i, page in enumerate(doc):
            ctx.check_cancel()
            ctx.progress(i / n, tr("Page {n}/{total}").format(n=i + 1, total=n))
            fname = safe_filename(tr("{name} page {n}").format(name=stem, n=f"{i + 1:0{pad}d}"))
            _save_pix(page.get_pixmap(dpi=dpi), folder / f"{fname}.{ext}", ext, quality)
        return folder


def _save_pix(pix, out: Path, ext: str, quality: int) -> None:
    if ext in ("jpg", "jpeg"):
        if pix.alpha:
            import pymupdf

            pix = pymupdf.Pixmap(pix, 0)
        pix.save(str(out), output="jpg", jpg_quality=quality)
    else:
        pix.save(str(out))


def pdf_to_jpg(src: Path, ctx) -> Path:
    return render_pages(src, "jpg", ctx)


def pdf_to_png(src: Path, ctx) -> Path:
    return render_pages(src, "png", ctx)


# --------------------------------------------------------------------------- office formats


def pdf_to_docx(src: Path, ctx) -> Path:
    from pdf2docx import Converter

    open_pdf(src).close()
    dst = ctx.output_path(src, "docx")
    ctx.progress(0.05, tr("Analysing layout…"))
    cv = Converter(str(src))
    try:
        cv.convert(str(dst))
    except Exception as e:
        raise UserError(tr("Could not convert {name} to Word: {error}").format(name=src.name, error=e)) from e
    finally:
        cv.close()
    return dst


EMU_PER_PT = 12700
_MIN_EMU, _MAX_EMU = 914400, 51206400  # PowerPoint slide size limits (1in .. 56in)


def pdf_to_pptx(src: Path, ctx) -> Path:
    """One slide per page, each holding the page rendered at 200 DPI."""
    from pptx import Presentation

    dst = ctx.output_path(src, "pptx")
    with open_pdf(src) as doc:
        first = doc[0].rect
        w, h = first.width * EMU_PER_PT, first.height * EMU_PER_PT
        scale = min(1.0, _MAX_EMU / max(w, h))
        if min(w, h) * scale < _MIN_EMU:
            scale = _MIN_EMU / min(w, h)
        sw, sh = int(w * scale), int(h * scale)
        prs = Presentation()
        prs.slide_width, prs.slide_height = sw, sh
        blank = prs.slide_layouts[6]
        n = doc.page_count
        for i, page in enumerate(doc):
            ctx.check_cancel()
            ctx.progress(i / n, tr("Page {n}/{total}").format(n=i + 1, total=n))
            png = page.get_pixmap(dpi=200).tobytes("png")
            slide = prs.slides.add_slide(blank)
            pw, ph = page.rect.width, page.rect.height
            fit = min(sw / pw, sh / ph)
            iw, ih = int(pw * fit), int(ph * fit)
            slide.shapes.add_picture(io.BytesIO(png), (sw - iw) // 2, (sh - ih) // 2, iw, ih)
        prs.save(str(dst))
    return dst


_NUM_RE = re.compile(r"^-?(0|[1-9]\d{0,2}(,\d{3})+|[1-9]\d*)(\.\d+)?$|^-?0?\.\d+$")


def cell_value(v):
    """Turn numeric-looking strings into numbers (keeps '007', dates, etc. as text)."""
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    if _NUM_RE.match(s):
        try:
            t = s.replace(",", "")
            return float(t) if "." in t else int(t)
        except ValueError:
            pass
    return s


def _clean_xml(s):
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    return ILLEGAL_CHARACTERS_RE.sub("", s) if isinstance(s, str) else s


def _sheet_title(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "_", name)[:31] or "Sheet"
    title, n = base, 2
    while title.lower() in used:
        tail = f" ({n})"
        title, n = base[:31 - len(tail)] + tail, n + 1
    used.add(title.lower())
    return title


def pdf_to_xlsx(src: Path, ctx) -> Path:
    """One sheet per detected table; falls back to the text lines of each page."""
    from openpyxl import Workbook

    dst = ctx.output_path(src, "xlsx")
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    with open_pdf(src) as doc:
        n = doc.page_count
        found = 0
        for pno, page in enumerate(doc, 1):
            ctx.check_cancel()
            ctx.progress((pno - 1) / n, tr("Page {n}/{total}").format(n=pno, total=n))
            tables = find_tables(page)
            for k, t in enumerate(tables, 1):
                rows = t.extract()
                if not rows or not any(any(c for c in r) for r in rows):
                    continue
                found += 1
                ws = wb.create_sheet(_sheet_title(tr("Page {page} Table {table}").format(page=pno, table=k), used))
                for r in rows:
                    ws.append([_clean_xml(cell_value(c)) for c in r])
        if not found:
            for pno, page in enumerate(doc, 1):
                lines = [ln for ln in page.get_text("text", sort=True).splitlines() if ln.strip()]
                ws = wb.create_sheet(_sheet_title(tr("Page {n}").format(n=pno), used))
                for ln in lines:
                    ws.append([_clean_xml(cell_value(c)) for c in re.split(r"\t| {2,}", ln.strip())])
    if not wb.worksheets:
        wb.create_sheet("Sheet1")
    wb.save(str(dst))
    return dst


# --------------------------------------------------------------------------- structured text

_BULLET_RE = re.compile(r"^\s*([•◦▪▫●○■□‣⁃∙·\-–\*])\s*(.*)$")
_NUMBERED_RE = re.compile(r"^\s*(\d{1,3}|[a-zA-Z])[.)]\s+(.*)$")


def _span_html(span) -> str:
    text = htmlmod.escape(span["text"])
    if not text.strip():
        return text
    bold = bool(span["flags"] & 16) or "bold" in span.get("font", "").lower()
    italic = bool(span["flags"] & 2) or "italic" in span.get("font", "").lower()
    if bold:
        text = f"<strong>{text}</strong>"
    if italic:
        text = f"<em>{text}</em>"
    return text


def _body_size(doc) -> float:
    sizes: Counter[float] = Counter()
    for i, page in enumerate(doc):
        if i >= 50:
            break
        for b in page.get_text("dict")["blocks"]:
            for line in b.get("lines", []):
                for s in line["spans"]:
                    sizes[round(s["size"], 1)] += len(s["text"].strip())
    return sizes.most_common(1)[0][0] if sizes else 11.0


def _inside(bbox, rects) -> bool:
    x0, y0, x1, y1 = bbox
    cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
    return any(r.x0 - 1 <= cx <= r.x1 + 1 and r.y0 - 1 <= cy <= r.y1 + 1 for r in rects)


def _table_html(rows) -> str:
    out = ["<table>"]
    for i, r in enumerate(rows):
        tag = "th" if i == 0 else "td"
        cells = "".join(f"<{tag}>{htmlmod.escape(str(c or '')).replace(chr(10), '<br/>')}</{tag}>"
                        for c in r)
        out.append(f"<tr>{cells}</tr>")
    out.append("</table>")
    return "".join(out)


def pdf_to_html_body(doc, ctx, media_dir: Path | None = None, embed: bool = False) -> str:
    """Rebuild headings (by font size), bold/italic, lists, tables and images as HTML."""
    import pymupdf

    body = _body_size(doc)
    big = sorted({round(s, 0) for s in _heading_sizes(doc, body)}, reverse=True)
    level = {s: min(i + 1, 4) for i, s in enumerate(big)}
    parts: list[str] = []
    img_no = 0
    n = doc.page_count
    for pno, page in enumerate(doc, 1):
        ctx.check_cancel()
        ctx.progress((pno - 1) / n, tr("Page {n}/{total}").format(n=pno, total=n))
        tables = find_tables(page)
        trects = [pymupdf.Rect(t.bbox) for t in tables]
        pending = sorted(((r.y0, _table_html(t.extract())) for r, t in zip(trects, tables)),
                         key=lambda x: x[0])
        items: list[str] = []
        for b in page.get_text("dict")["blocks"]:
            y0 = b["bbox"][1]
            while pending and pending[0][0] <= y0:
                items.append(pending.pop(0)[1])
            if b["type"] == 1:
                if b.get("width", 0) < _IMG_MIN or b.get("height", 0) < _IMG_MIN or not b.get("image"):
                    continue
                img_no += 1
                ext = b.get("ext", "png")
                if embed or media_dir is None:
                    mime = "jpeg" if ext in ("jpg", "jpeg") else ext
                    src = f"data:image/{mime};base64,{base64.b64encode(b['image']).decode()}"
                else:
                    media_dir.mkdir(parents=True, exist_ok=True)
                    name = f"image{img_no:03d}.{ext}"
                    (media_dir / name).write_bytes(b["image"])
                    src = f"{media_dir.name}/{name}"
                alt = htmlmod.escape(tr("Image {n}").format(n=img_no))
                items.append(f'<p><img src="{src}" alt="{alt}"/></p>')
                continue
            if _inside(b["bbox"], trects):
                continue
            items.extend(_text_block_html(b, body, level))
        items.extend(h for _, h in pending)
        parts.append("\n".join(items))
    return "\n".join(parts)


def _heading_sizes(doc, body: float) -> set[float]:
    out: set[float] = set()
    for i, page in enumerate(doc):
        if i >= 50:
            break
        for b in page.get_text("dict")["blocks"]:
            for line in b.get("lines", []):
                for s in line["spans"]:
                    if s["text"].strip() and s["size"] >= body * 1.15:
                        out.add(s["size"])
    return out


def _spans_html(spans, skip: int = 0) -> str:
    """HTML for a line's spans, dropping the first `skip` characters (list markers)."""
    out = []
    for s in spans:
        text = s["text"]
        if skip:
            cut = min(skip, len(text))
            text, skip = text[cut:], skip - cut
            if not text:
                continue
        out.append(_span_html({**s, "text": text}))
    return "".join(out).strip()


def _text_block_html(block, body: float, level: dict[float, int]) -> list[str]:
    lines = []
    for line in block.get("lines", []):
        spans = [s for s in line["spans"] if s["text"]]
        plain = "".join(s["text"] for s in spans)
        if not plain.strip():
            continue
        size = max(s["size"] for s in spans if s["text"].strip())
        lines.append((plain, spans, size))
    if not lines:
        return []
    text_len = sum(len(p) for p, _, _ in lines)
    if all(sz >= body * 1.15 for _, _, sz in lines) and text_len < 300:
        lv = level.get(round(max(sz for _, _, sz in lines), 0), 4)
        inner = " ".join(htmlmod.escape(p.strip()) for p, _, _ in lines)
        return [f"<h{lv}>{inner}</h{lv}>"]
    out: list[str] = []
    para: list[str] = []
    lst: list[str] = []
    ordered = False

    def flush_para():
        if para:
            out.append("<p>" + _join(para) + "</p>")
            para.clear()

    def flush_list():
        if lst:
            tag = "ol" if ordered else "ul"
            out.append(f"<{tag}>" + "".join(f"<li>{x}</li>" for x in lst) + f"</{tag}>")
            lst.clear()

    for plain, spans, _ in lines:
        mb, mn = _BULLET_RE.match(plain), _NUMBERED_RE.match(plain)
        m = mb if (mb and mb.group(2).strip()) else mn
        if m:
            flush_para()
            is_num = m is mn
            if lst and is_num != ordered:
                flush_list()
            ordered = is_num
            lst.append(_spans_html(spans, skip=m.start(2)))
        elif lst:
            lst[-1] = _join([lst[-1], _spans_html(spans)])
        else:
            para.append(_spans_html(spans))
    flush_para()
    flush_list()
    return out


def _join(chunks: list[str]) -> str:
    out = ""
    for c in chunks:
        c = c.strip()
        if not out:
            out = c
        elif out.endswith("-") and c[:1].islower():
            out = out[:-1] + c
        else:
            out += " " + c
    return out.replace("</strong> <strong>", " ").replace("</em> <em>", " ")


def pdf_to_txt(src: Path, ctx) -> Path:
    dst = ctx.output_path(src, "txt")
    with open_pdf(src) as doc:
        pages = []
        for i, page in enumerate(doc):
            ctx.check_cancel()
            pages.append(page.get_text("text", sort=True).rstrip())
    dst.write_text("\n\n".join(pages) + "\n", encoding="utf-8")
    return dst


def pdf_to_html(src: Path, ctx) -> Path:
    """A single self-contained HTML file (images embedded)."""
    from .documents import BASE_CSS

    dst = ctx.output_path(src, "html")
    with open_pdf(src) as doc:
        body = pdf_to_html_body(doc, ctx, embed=True)
    css = BASE_CSS + "body { max-width: 50em; margin: 2em auto; padding: 0 1em; }"
    dst.write_text(wrap_html(body, src.stem, css), encoding="utf-8")
    return dst


def _via_pandoc(src: Path, ctx, target: str) -> Path:
    work = ctx.tmpdir / "pdfhtml"
    work.mkdir(exist_ok=True)
    with open_pdf(src) as doc:
        body = pdf_to_html_body(doc, ctx, media_dir=work / "media")
    (work / "in.html").write_text(wrap_html(body, src.stem), encoding="utf-8")
    dst = ctx.output_path(src, target)
    ctx.progress(0.9, tr("Writing…"))
    if target == "md":
        args = ["-f", "html", "-t", "gfm", "--wrap=none"]
        if (work / "media").is_dir():
            args.append(f"--extract-media={media_dir_name(dst)}")
        cwd = dst.parent
    else:
        args = ["-f", "html", "-t", "epub3", "--metadata", f"title={src.stem}"]
        cwd = work
    pandoc([*args, f"--resource-path={work}", "-o", str(dst.resolve()), str((work / "in.html").resolve())],
           ctx, cwd=cwd)
    return dst


def pdf_to_md(src: Path, ctx) -> Path:
    return _via_pandoc(src, ctx, "md")


def pdf_to_epub(src: Path, ctx) -> Path:
    return _via_pandoc(src, ctx, "epub")


# --------------------------------------------------------------------------- registration

_RUNNERS = {"docx": (pdf_to_docx, "DOCX"), "pptx": (pdf_to_pptx, "PPTX"), "xlsx": (pdf_to_xlsx, "XLSX"),
            "jpg": (pdf_to_jpg, "JPG"), "png": (pdf_to_png, "PNG"), "epub": (pdf_to_epub, "EPUB"),
            "txt": (pdf_to_txt, "TXT"), "md": (pdf_to_md, "Markdown"), "html": (pdf_to_html, "HTML")}

for _i, _target in enumerate(formats.PDF_OUT):
    _fn, _label = _RUNNERS[_target]
    add(Option(id=f"pdf.to.{_target}", label=_label, mode="convert", run=_fn, exts=frozenset(formats.PDF_IN),
               family="pdf", order=20 + _i, skip_same_ext=_target,
               description="Convert PDF to {fmt}", desc_args={"fmt": _label}))
