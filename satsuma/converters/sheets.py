"""Spreadsheets: XLSX XLS XLSM ODS CSV TSV -> XLSX ODS CSV TSV PDF HTML."""
from __future__ import annotations

import csv
import datetime as dt
import html as htmlmod
import io
import math
import re
from pathlib import Path

from .. import formats, office
from ..i18n import tr
from ..paths import localized_stem
from ..registry import Option, UserError, add

Sheet = tuple[str, list[list]]
WORKBOOK_IN = {"xlsx", "xls", "xlsm", "ods"}


# --------------------------------------------------------------------------- reading


def decode_csv(path: Path) -> str:
    data = Path(path).read_bytes()
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def sniff_delimiter(text: str, default: str = ",") -> str:
    sample = "\n".join(text.splitlines()[:50])
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        counts = {d: sample.count(d) for d in ",;\t|"}
        best = max(counts, key=counts.get)
        return best if counts[best] else default


_PLAIN_NUM = re.compile(r"^-?(0|[1-9]\d*)(\.\d+)?$")


def _csv_value(s: str):
    s2 = s.strip()
    if _PLAIN_NUM.match(s2) and len(s2) < 16:
        return float(s2) if "." in s2 else int(s2)
    return s if s != "" else None


def _trim(rows: list[list]) -> list[list]:
    out = []
    for r in rows:
        r = list(r)
        while r and (r[-1] is None or r[-1] == ""):
            r.pop()
        out.append(r)
    while out and not out[-1]:
        out.pop()
    return out


def read_sheets(src: Path) -> list[Sheet]:
    """Every sheet of `src` as (name, rows of python values)."""
    kind = formats.ext(src)
    try:
        if kind in ("csv", "tsv"):
            text = decode_csv(src)
            delim = "\t" if kind == "tsv" else sniff_delimiter(text)
            rows = [[_csv_value(c) for c in r] for r in csv.reader(io.StringIO(text), delimiter=delim)]
            return [(src.stem[:31] or "Sheet1", _trim(rows))]
        if kind in ("xlsx", "xlsm"):
            from openpyxl import load_workbook

            wb = load_workbook(str(src), read_only=True, data_only=True)
            try:
                return [(ws.title, _trim([list(r) for r in ws.iter_rows(values_only=True)]))
                        for ws in wb.worksheets]
            finally:
                wb.close()
        if kind == "xls":
            import xlrd

            book = xlrd.open_workbook(str(src))
            sheets = []
            for sh in book.sheets():
                rows = []
                for r in range(sh.nrows):
                    row = []
                    for c in range(sh.ncols):
                        cell = sh.cell(r, c)
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                        elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            row.append(bool(cell.value))
                        elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                            row.append(None)
                        else:
                            row.append(cell.value)
                    rows.append(row)
                sheets.append((sh.name, _trim(rows)))
            return sheets
        if kind == "ods":
            import pandas as pd

            frames = pd.read_excel(str(src), sheet_name=None, header=None, engine="odf")
            return [(name, _trim([[_from_pandas(v) for v in row] for row in df.itertuples(index=False)]))
                    for name, df in frames.items()]
    except UserError:
        raise
    except Exception as e:
        raise UserError(tr("Cannot read {name}: {error}").format(name=src.name, error=e)) from e
    raise UserError(tr("Unsupported spreadsheet: {name}").format(name=src.name))


def _from_pandas(v):
    if v is None:
        return None
    if isinstance(v, float) and math.isnan(v):
        return None
    if hasattr(v, "to_pydatetime"):
        try:
            return v.to_pydatetime()
        except Exception:
            return None
    if hasattr(v, "item"):  # numpy scalar
        return v.item()
    return v


# --------------------------------------------------------------------------- writing


def fmt_cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        if math.isnan(v):
            return ""
        return str(int(v)) if v.is_integer() and abs(v) < 1e15 else repr(v)
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat(sep=" ")
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    return str(v)


def _safe_title(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "_", str(name))[:31] or "Sheet"
    title, n = base, 2
    while title.lower() in used:
        tail = f" ({n})"
        title, n = base[:31 - len(tail)] + tail, n + 1
    used.add(title.lower())
    return title


def write_xlsx(sheets: list[Sheet], dst: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for name, rows in sheets:
        ws = wb.create_sheet(_safe_title(name, used))
        for r in rows:
            ws.append([ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in r])
    if not wb.worksheets:
        wb.create_sheet("Sheet1")
    wb.save(str(dst))
    return dst


def write_ods(sheets: list[Sheet], dst: Path) -> Path:
    import pandas as pd

    used: set[str] = set()
    with pd.ExcelWriter(str(dst), engine="odf") as xw:
        for name, rows in sheets or [("Sheet1", [])]:
            width = max((len(r) for r in rows), default=0)
            df = pd.DataFrame([r + [None] * (width - len(r)) for r in rows])
            df.to_excel(xw, sheet_name=_safe_title(name, used), header=False, index=False)
    return dst


def write_delimited(rows: list[list], dst: Path, delim: str) -> Path:
    with open(dst, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=delim)
        for r in rows:
            w.writerow([fmt_cell(v) for v in r])
    return dst


def sheets_html(sheets: list[Sheet], title: str, headings: bool = True) -> str:
    parts = []
    for name, rows in sheets:
        if headings and len(sheets) > 1:
            parts.append(f"<h2>{htmlmod.escape(str(name))}</h2>")
        width = max((len(r) for r in rows), default=0)
        cells = []
        for i, r in enumerate(rows):
            tag = "th" if i == 0 else "td"
            vals = list(r) + [None] * (width - len(r))
            cells.append("<tr>" + "".join(f"<{tag}>{htmlmod.escape(fmt_cell(v))}</{tag}>" for v in vals)
                         + "</tr>")
        parts.append("<table>" + "".join(cells) + "</table>" if cells
                     else f"<p>{htmlmod.escape(tr('(empty)'))}</p>")
    return "\n".join(parts)


HTML_CSS = """
body { font-family: "Segoe UI", sans-serif; font-size: 10pt; margin: 2em; }
table { border-collapse: collapse; margin-bottom: 2em; }
th, td { border: 1px solid #bbb; padding: 3px 6px; text-align: left; vertical-align: top; }
th { background: #f0f0f0; }
"""


# --------------------------------------------------------------------------- conversion


def convert_sheet(src: Path, target: str, ctx) -> Path:
    from .documents import BASE_CSS, html_to_pdf, wrap_html

    kind = formats.ext(src)
    office_ok = office.excel_available() or office.libreoffice_available()

    # Workbook -> workbook / PDF: Excel keeps formatting, formulas and layout.
    if office_ok and kind in WORKBOOK_IN and target in ("xlsx", "ods", "pdf"):
        dst = ctx.output_path(src, target)
        ctx.progress(0.1, tr("Converting with Excel…") if office.excel_available() else tr("Converting…"))
        try:
            return office.convert_with_office("excel", src, dst, target, ctx)
        except office.OfficeError:
            pass  # fall back to the pure-python path below

    ctx.progress(0.2, tr("Reading…"))
    sheets = read_sheets(src)
    ctx.check_cancel()

    if target == "xlsx":
        return write_xlsx(sheets, ctx.output_path(src, "xlsx"))
    if target == "ods":
        return write_ods(sheets, ctx.output_path(src, "ods"))
    if target in ("csv", "tsv"):
        delim = "\t" if target == "tsv" else ","
        nonempty = sheets if len(sheets) <= 1 else [s for s in sheets if s[1]] or sheets[:1]
        if len(nonempty) <= 1:
            rows = nonempty[0][1] if nonempty else []
            return write_delimited(rows, ctx.output_path(src, target), delim)
        folder = ctx.output_path(src, "", name=localized_stem(tr("{name} sheets"), src))
        folder.mkdir(parents=True)
        used: set[str] = set()
        for name, rows in nonempty:
            fname = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(name)).strip(" .") or "Sheet"
            while fname.lower() in used:
                fname += "_"
            used.add(fname.lower())
            write_delimited(rows, folder / f"{fname}.{target}", delim)
        return folder
    if target == "html":
        dst = ctx.output_path(src, "html")
        dst.write_text(wrap_html(sheets_html(sheets, src.stem), src.stem, HTML_CSS), encoding="utf-8")
        return dst
    if target == "pdf":
        dst = ctx.output_path(src, "pdf")
        if office.excel_available() and kind in ("csv", "tsv"):
            tmp = write_xlsx(sheets, ctx.tmpdir / f"{src.stem}.xlsx")
            try:
                return office.excel_convert(tmp, dst, "pdf")
            except office.OfficeError:
                pass
        width = max((len(r) for _, rows in sheets for r in rows), default=0)
        css = BASE_CSS + "th, td { font-size: 8pt; }"
        html_to_pdf(wrap_html(sheets_html(sheets, src.stem), src.stem), dst, landscape=width > 6,
                    css=css, ctx=ctx)
        return dst
    raise UserError(tr("Cannot convert to {fmt}").format(fmt=target.upper()))


def _make(target: str):
    def run(src: Path, ctx):
        return convert_sheet(Path(src), target, ctx)

    run.__name__ = f"sheet_to_{target}"
    return run


for _i, _target in enumerate(formats.SHEET_OUT):
    add(Option(id=f"sheet.to.{_target}", label=_target.upper(), mode="convert", run=_make(_target),
               exts=frozenset(formats.SHEET_IN), family="sheet", order=20 + _i, skip_same_ext=_target,
               description="Convert spreadsheets to {fmt}", desc_args={"fmt": _target.upper()}))
