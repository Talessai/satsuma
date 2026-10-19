"""Tests for PDF / document / slides / sheet / subtitle / archive converters and PDF tools.

Sample inputs are generated on the fly. Office COM tests run when Microsoft Office is
installed and skip (not fail) when automation is unavailable.
"""
from __future__ import annotations

import csv
import io
import os
import shutil
import subprocess
import tarfile
import zipfile
from pathlib import Path

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pymupdf  # noqa: E402

from satsuma import office, registry  # noqa: E402
from satsuma.config import settings  # noqa: E402
from satsuma.registry import UserError  # noqa: E402
from tests.helpers import FakeCtx, run_option  # noqa: E402

registry.load_all()


# --------------------------------------------------------------------------- fixtures


def _png_bytes(w=240, h=160, color=(30, 120, 220)) -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, w, h), 0)
    pix.set_rect(pix.irect, color)
    pix.set_rect(pymupdf.IRect(0, 0, w // 2, h // 2), (250, 200, 40))
    return pix.tobytes("png")


def _noise_jpeg(w=1600, h=1200) -> bytes:
    """A large, detailed photo-like image so compression has something to do."""
    from PIL import Image

    img = Image.effect_noise((w, h), 80).convert("RGB")
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=95)
    return buf.getvalue()


def make_pdf(path: Path) -> Path:
    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((72, 80), "Quarterly Report", fontsize=24, fontname="hebo")
    p.insert_text((72, 120), "This report summarises the results for the quarter.", fontsize=11)
    p.insert_text((72, 140), "Important:", fontsize=11, fontname="hebo")
    p.insert_text((72, 170), "• Apples grew strongly", fontsize=11)
    p.insert_text((72, 186), "• Pears were stable", fontsize=11)
    # a ruled 3x3 table
    x0, y0, cw, rh = 72, 220, 120, 22
    cells = [["Fruit", "Q1", "Q2"], ["Apples", "10", "12"], ["Pears", "7", "8"]]
    for r in range(4):
        p.draw_line((x0, y0 + r * rh), (x0 + 3 * cw, y0 + r * rh))
    for c in range(4):
        p.draw_line((x0 + c * cw, y0), (x0 + c * cw, y0 + 3 * rh))
    for r, row in enumerate(cells):
        for c, val in enumerate(row):
            p.insert_text((x0 + c * cw + 5, y0 + r * rh + 15), val, fontsize=10)
    p2 = doc.new_page()
    p2.insert_text((72, 80), "Pictures", fontsize=24, fontname="hebo")
    p2.insert_image(pymupdf.Rect(72, 110, 312, 270), stream=_png_bytes())
    p2.insert_text((72, 300), "A picture is worth a thousand words.", fontsize=11)
    p3 = doc.new_page()
    p3.insert_text((72, 80), "Final page text", fontsize=11)
    doc.save(str(path))
    return path


@pytest.fixture(scope="session")
def S(tmp_path_factory):
    """All sample inputs, generated once."""
    settings.set("output_location", "same")
    settings.set("image_dpi", 72)
    d = tmp_path_factory.mktemp("samples")
    s: dict[str, Path] = {"dir": d}

    s["pdf"] = make_pdf(d / "report.pdf")

    big = pymupdf.open()
    for _ in range(2):
        pg = big.new_page()
        pg.insert_image(pg.rect, stream=_noise_jpeg())
        pg.insert_text((72, 72), "Photo page", fontsize=20)
    big.save(str(d / "photos.pdf"))
    s["bigpdf"] = d / "photos.pdf"

    import docx

    document = docx.Document()
    document.add_heading("Project Plan", level=1)
    document.add_paragraph("An introduction paragraph with ").add_run("bold text").bold = True
    document.add_paragraph("First bullet", style="List Bullet")
    document.add_paragraph("Second bullet", style="List Bullet")
    table = document.add_table(rows=2, cols=3)
    for r in range(2):
        for c in range(3):
            table.cell(r, c).text = f"R{r}C{c}"
    document.add_picture(io.BytesIO(_png_bytes()))
    document.save(str(d / "plan.docx"))
    s["docx"] = d / "plan.docx"

    s["md"] = d / "notes.md"
    s["md"].write_text("# Meeting Notes\n\nSome **bold** words.\n\n- one\n- two\n\n| a | b |\n|---|---|\n"
                       "| 1 | 2 |\n", encoding="utf-8")
    s["txt"] = d / "plain.txt"
    s["txt"].write_text("Hello plain text.\nSecond line.\n\nNew paragraph with <angle> & ampersand.\n",
                        encoding="utf-8")
    s["html"] = d / "page.html"
    s["html"].write_text("<html><head><title>T</title></head><body><h1>Web Page</h1><p>Para "
                         "<b>strong</b></p><ul><li>x</li></ul></body></html>", encoding="utf-8")

    rtfd = d / "Letter.rtfd"
    rtfd.mkdir()
    (rtfd / "TXT.rtf").write_text(r"{\rtf1\ansi\deff0{\fonttbl{\f0 Helvetica;}}\f0\fs24 Dear reader,"
                                  r"\par This is {\b RTFD} content.\par}", encoding="ascii")
    s["rtfd"] = rtfd

    pages = d / "Brochure.pages"
    with zipfile.ZipFile(pages, "w") as z:
        z.writestr("Index/Document.iwa", b"\x00")
        z.writestr("QuickLook/Preview.pdf", s["pdf"].read_bytes())
    s["pages"] = pages
    nopreview = d / "Empty.pages"
    with zipfile.ZipFile(nopreview, "w") as z:
        z.writestr("Index/Document.iwa", b"\x00")
    s["pages_nopreview"] = nopreview

    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for title in ("Welcome slide", "Second slide"):
        sl = prs.slides.add_slide(prs.slide_layouts[1])
        sl.shapes.title.text = title
        sl.placeholders[1].text = f"Body of {title}"
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.save(str(d / "deck.pptx"))
    s["pptx"] = d / "deck.pptx"

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Amount", "Code"])
    ws.append(["Alice", 10.5, "007"])
    ws.append(["Bob", 3, "042"])
    ws2 = wb.create_sheet("Notes")
    ws2.append(["Remember the milk"])
    wb.save(str(d / "book.xlsx"))
    s["xlsx"] = d / "book.xlsx"

    s["csv"] = d / "semi.csv"
    s["csv"].write_bytes("name;price;city\nTea;2.5;Café Town\nCoffee;3;Zürich\n".encode("cp1252"))

    s["srt"] = d / "movie.srt"
    s["srt"].write_text("1\n00:00:01,000 --> 00:00:02,500\n<i>Hello there.</i>\n\n"
                        "2\n00:00:03,000 --> 00:00:04,000\nGeneral Kenobi!\nYou are a bold one.\n\n"
                        "3\n01:00:05,250 --> 01:00:06,000\nBye\n", encoding="utf-8")
    s["vtt"] = d / "clip.vtt"
    s["vtt"].write_text("WEBVTT\n\nNOTE a comment\n\nintro\n00:01.000 --> 00:02.000 align:start\n"
                        "<v Roger>Hi <b>you</b></v>\n\n00:00:02.000 --> 00:00:03.500\nSecond cue\n",
                        encoding="utf-8")

    z = d / "bundle.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("docs/readme.txt", "inside zip")
        zf.writestr("top.txt", "top level")
        zf.writestr("../evil.txt", "escaped!")
        zf.writestr("/abs/evil2.txt", "escaped!")
    s["zip"] = z

    t = d / "bundle.tar"
    with tarfile.open(t, "w") as tf:
        for name, data in (("folder/a.txt", b"tar a"), ("folder/b.txt", b"tar b"), ("../evil.txt", b"x")):
            info = tarfile.TarInfo(name)
            info.size = len(data)
            tf.addfile(info, io.BytesIO(data))
    s["tar"] = t

    tg = d / "single.tar.gz"
    with tarfile.open(tg, "w:gz") as tf:
        info = tarfile.TarInfo("only.txt")
        info.size = 4
        tf.addfile(info, io.BytesIO(b"only"))
    s["tgz"] = tg

    folder = d / "My Folder"
    (folder / "sub").mkdir(parents=True)
    (folder / "one.txt").write_text("1")
    (folder / "sub" / "two.txt").write_text("2")
    s["folder"] = folder
    return s


def office_call(fn, *a, **kw):
    """Run something that needs Office COM; skip the test if Office automation fails."""
    try:
        return fn(*a, **kw)
    except (office.OfficeError, UserError) as e:
        msg = str(e)
        if any(k in msg for k in ("Office", "Word", "Excel", "PowerPoint", "LibreOffice", "COM")):
            pytest.skip(f"Office automation unavailable: {msg}")
        raise


def pdf_text(path: Path) -> str:
    with pymupdf.open(str(path)) as d:
        return "".join(p.get_text() for p in d)


def pdf_pages(path: Path) -> int:
    with pymupdf.open(str(path)) as d:
        return d.page_count


@pytest.fixture
def no_office(monkeypatch):
    """Force the pure-Python fallbacks."""
    monkeypatch.setattr(office, "word_available", lambda: False)
    monkeypatch.setattr(office, "excel_available", lambda: False)
    monkeypatch.setattr(office, "powerpoint_available", lambda: False)
    monkeypatch.setattr(office, "libreoffice_available", lambda: False)
    import satsuma.paths

    monkeypatch.setattr(satsuma.paths, "soffice_exe", lambda: None)
    monkeypatch.setattr(office, "soffice_exe", lambda: None)


# --------------------------------------------------------------------------- registry


def test_options_registered_and_apply(S):
    ids = {o.id for o in registry.all_options()}
    for i in ["pdf.to." + t for t in ("docx", "pptx", "xlsx", "jpg", "png", "epub", "txt", "md", "html")]:
        assert i in ids
    conv = {o.id for o in registry.options_for("convert", [S["docx"]])}
    assert "doc.to.pdf" in conv and "doc.to.docx" not in conv
    assert "doc.to.rtf" in {o.id for o in registry.options_for("convert", [S["rtfd"]])}
    assert "doc.to.pdf" not in {o.id for o in registry.options_for("convert", [S["folder"]])}
    assert "any.to.zip" in {o.id for o in registry.options_for("convert", [S["folder"]])}
    assert {"pages.to.pdf", "pages.to.png"} <= {o.id for o in registry.options_for("convert", [S["pages"]])}
    zconv = {o.id for o in registry.options_for("convert", [S["zip"]])}
    assert "archive.extract" in zconv and "archive.to.zip" not in zconv and "any.to.zip" not in zconv
    tools = {o.id for o in registry.options_for("tools", [S["pdf"], S["bigpdf"]])}
    assert "pdf.merge" in tools
    assert "pdf.merge" not in {o.id for o in registry.options_for("tools", [S["pdf"]])}


# --------------------------------------------------------------------------- PDF converters


def test_pdf_to_docx(S):
    import docx

    (out,) = run_option("pdf.to.docx", [S["pdf"]])
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "Quarterly Report" in text


def test_pdf_to_pptx(S):
    from pptx import Presentation

    (out,) = run_option("pdf.to.pptx", [S["pdf"]])
    prs = Presentation(str(out))
    assert len(prs.slides) == 3
    assert abs(prs.slide_width / prs.slide_height - 595 / 842) < 0.01


def test_pdf_to_xlsx(S):
    from openpyxl import load_workbook

    (out,) = run_option("pdf.to.xlsx", [S["pdf"]])
    wb = load_workbook(str(out))
    values = [c for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row]
    assert "Apples" in values and 12 in values


def test_pdf_to_images(S):
    (folder,) = run_option("pdf.to.png", [S["pdf"]])
    assert folder.is_dir() and folder.name == "report pages"
    files = sorted(folder.iterdir())
    assert [f.name for f in files] == [f"report page 0{i}.png" for i in (1, 2, 3)]
    pix = pymupdf.Pixmap(str(files[0]))
    assert pix.width == 595  # 72 DPI in tests
    single = S["dir"] / "one.pdf"
    with pymupdf.open(str(S["pdf"])) as d:
        d.select([0])
        d.save(str(single))
    (jpg,) = run_option("pdf.to.jpg", [single])
    assert jpg.suffix == ".jpg" and jpg.is_file()


def test_pdf_to_txt_md_html(S):
    (txt,) = run_option("pdf.to.txt", [S["pdf"]])
    t = txt.read_text(encoding="utf-8")
    assert "Quarterly Report" in t and "Final page text" in t
    (md,) = run_option("pdf.to.md", [S["pdf"]])
    m = md.read_text(encoding="utf-8")
    assert "# Quarterly Report" in m
    assert "**Important:**" in m
    assert "- Apples grew strongly" in m
    assert "Apples" in m and "|" in m  # table
    assert (md.parent / "report_media").is_dir()  # extracted image
    (html,) = run_option("pdf.to.html", [S["pdf"]])
    h = html.read_text(encoding="utf-8")
    assert "<h1>Quarterly Report</h1>" in h and "<table>" in h and "data:image/png;base64" in h
    assert "<li>Pears were stable</li>" in h


def test_pdf_to_epub(S):
    (epub,) = run_option("pdf.to.epub", [S["pdf"]])
    from ebooklib import epub as ebl

    book = ebl.read_epub(str(epub))
    content = b"".join(i.get_content() for i in book.get_items() if i.media_type == "application/xhtml+xml")
    assert b"Quarterly Report" in content
    assert any(i.media_type.startswith("image/") for i in book.get_items())


def test_pdf_locked_gives_user_error(S):
    locked = S["dir"] / "locked.pdf"
    with pymupdf.open(str(S["pdf"])) as d:
        d.save(str(locked), encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="pw", owner_pw="o")
    with pytest.raises(UserError, match="password"):
        run_option("pdf.to.txt", [locked])


# --------------------------------------------------------------------------- documents


def test_docx_to_text_formats(S):
    (md,) = run_option("doc.to.md", [S["docx"]])
    m = md.read_text(encoding="utf-8")
    assert "# Project Plan" in m and "**bold text**" in m and "First bullet" in m
    (txt,) = run_option("doc.to.txt", [S["docx"]])
    assert "Project Plan" in txt.read_text(encoding="utf-8")
    (html,) = run_option("doc.to.html", [S["docx"]])
    h = html.read_text(encoding="utf-8")
    assert "Project Plan" in h and "data:image/png;base64" in h
    (epub,) = run_option("doc.to.epub", [S["docx"]])
    with pymupdf.open(str(epub)) as book:
        assert "Project Plan" in "".join(p.get_text() for p in book)


def test_docx_to_office_formats(S):
    from odf.opendocument import load as odf_load

    (odt,) = office_call(run_option, "doc.to.odt", [S["docx"]])
    assert odt.suffix == ".odt"
    odf_load(str(odt))
    (rtf,) = office_call(run_option, "doc.to.rtf", [S["docx"]])
    assert rtf.read_bytes().startswith(b"{\\rtf")


def test_docx_to_pdf_word(S):
    if not office.word_available():
        pytest.skip("Word not installed")
    (pdf,) = office_call(run_option, "doc.to.pdf", [S["docx"]])
    assert "Project Plan" in pdf_text(pdf)


def test_docx_to_pdf_fallback(S, no_office):
    (pdf,) = run_option("doc.to.pdf", [S["docx"]])
    t = pdf_text(pdf)
    assert "Project Plan" in t and "R1C2" in t


def test_docx_to_odt_rtf_fallback(S, no_office):
    (odt,) = run_option("doc.to.odt", [S["docx"]])
    assert zipfile.is_zipfile(odt)
    (rtf,) = run_option("doc.to.rtf", [S["docx"]])
    assert b"Project Plan" in rtf.read_bytes()


def test_md_txt_html_conversions(S):
    import docx

    (d1,) = run_option("doc.to.docx", [S["md"]])
    doc = docx.Document(str(d1))
    assert doc.paragraphs[0].text == "Meeting Notes" and len(doc.tables) == 1
    (p1,) = run_option("doc.to.pdf", [S["md"]])
    assert "Meeting Notes" in pdf_text(p1)
    (h1,) = run_option("doc.to.html", [S["md"]])
    assert "<strong>bold</strong>" in h1.read_text(encoding="utf-8")
    (p2,) = run_option("doc.to.pdf", [S["txt"]])
    assert "<angle> & ampersand" in pdf_text(p2)
    (d2,) = run_option("doc.to.docx", [S["txt"]])
    assert any("Second line" in p.text for p in docx.Document(str(d2)).paragraphs)
    (m2,) = run_option("doc.to.md", [S["html"]])
    assert "# Web Page" in m2.read_text(encoding="utf-8")
    (t2,) = run_option("doc.to.txt", [S["md"]])
    assert "Meeting Notes" in t2.read_text(encoding="utf-8")


def test_epub_to_pdf(S):
    (epub,) = run_option("doc.to.epub", [S["md"]])
    (pdf,) = run_option("doc.to.pdf", [epub])
    assert "Meeting Notes" in pdf_text(pdf)


def test_rtfd(S, no_office):
    (rtf,) = run_option("doc.to.rtf", [S["rtfd"]])
    assert rtf.name == "Letter.rtf" and b"RTFD" in rtf.read_bytes()
    (md,) = run_option("doc.to.md", [S["rtfd"]])
    assert md.name == "Letter.md" and "RTFD" in md.read_text(encoding="utf-8")


def test_pages_preview(S):
    (pdf,) = run_option("pages.to.pdf", [S["pages"]])
    assert pdf_pages(pdf) == 3
    (png,) = run_option("pages.to.png", [S["pages"]])
    assert png.is_dir() and len(list(png.glob("*.png"))) == 3
    with pytest.raises(UserError, match="preview"):
        run_option("pages.to.pdf", [S["pages_nopreview"]])


def test_doc_via_word(S):
    """Legacy .doc: made with Word, converted back through Word."""
    if not office.word_available():
        pytest.skip("Word not installed")
    doc = S["dir"] / "legacy.doc"
    office_call(office.word_convert, S["docx"], doc, "doc")
    (md,) = office_call(run_option, "doc.to.md", [doc])
    assert "Project Plan" in md.read_text(encoding="utf-8")


# --------------------------------------------------------------------------- slides


@pytest.fixture
def quiet_faulthandler():
    """PowerPoint exits as soon as it is told to quit; releasing the proxy afterwards
    raises a harmless, handled RPC_E_DISCONNECTED that faulthandler would print."""
    import faulthandler

    was = faulthandler.is_enabled()
    faulthandler.disable()
    yield
    if was:
        faulthandler.enable()


def test_slides_pdf_and_images(S, quiet_faulthandler):
    if not office.powerpoint_available():
        pytest.skip("PowerPoint not installed")
    (pdf,) = office_call(run_option, "slides.to.pdf", [S["pptx"]])
    assert pdf_pages(pdf) == 2 and "Welcome slide" in pdf_text(pdf)
    (folder,) = office_call(run_option, "slides.to.png", [S["pptx"]])
    files = sorted(folder.glob("*.png"))
    assert folder.name == "deck slides" and len(files) == 2
    assert pymupdf.Pixmap(str(files[0])).width == 1920


def test_slides_need_office(S, no_office):
    with pytest.raises(UserError, match="PowerPoint or LibreOffice"):
        run_option("slides.to.pdf", [S["pptx"]])


# --------------------------------------------------------------------------- sheets


def test_xlsx_to_csv_tsv_html(S):
    (folder,) = run_option("sheet.to.csv", [S["xlsx"]])
    assert folder.is_dir()
    data = (folder / "Data.csv").read_text(encoding="utf-8-sig").splitlines()
    assert data[0] == "Name,Amount,Code" and data[1] == "Alice,10.5,007" and data[2] == "Bob,3,042"
    assert (folder / "Notes.csv").exists()
    (tfolder,) = run_option("sheet.to.tsv", [S["xlsx"]])
    assert "Alice\t10.5\t007" in (tfolder / "Data.tsv").read_text(encoding="utf-8-sig")
    (html,) = run_option("sheet.to.html", [S["xlsx"]])
    h = html.read_text(encoding="utf-8")
    assert "<h2>Notes</h2>" in h and "<td>Alice</td>" in h


def test_xlsx_to_ods_and_pdf_fallback(S, no_office):
    import pandas as pd

    (ods,) = run_option("sheet.to.ods", [S["xlsx"]])
    frames = pd.read_excel(str(ods), sheet_name=None, header=None, engine="odf")
    assert list(frames) == ["Data", "Notes"]
    (pdf,) = run_option("sheet.to.pdf", [S["xlsx"]])
    assert "Alice" in pdf_text(pdf)


def test_xlsx_office(S):
    if not office.excel_available():
        pytest.skip("Excel not installed")
    (pdf,) = office_call(run_option, "sheet.to.pdf", [S["xlsx"]])
    assert "Alice" in pdf_text(pdf)
    (ods,) = office_call(run_option, "sheet.to.ods", [S["xlsx"]])
    import pandas as pd

    assert len(pd.read_excel(str(ods), sheet_name=None, engine="odf")) == 2


def test_csv_sniffing_and_encoding(S):
    from openpyxl import load_workbook

    (xlsx,) = run_option("sheet.to.xlsx", [S["csv"]])
    ws = load_workbook(str(xlsx)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("name", "price", "city")
    assert rows[1] == ("Tea", 2.5, "Café Town") and rows[2][1] == 3
    (tsv,) = run_option("sheet.to.tsv", [S["csv"]])
    assert tsv.is_file()
    rows = list(csv.reader(tsv.read_text(encoding="utf-8-sig").splitlines(), delimiter="\t"))
    assert rows[2] == ["Coffee", "3", "Zürich"]


def test_csv_to_pdf_fallback(S, no_office):
    (pdf,) = run_option("sheet.to.pdf", [S["csv"]])
    assert "Zürich" in pdf_text(pdf)


# --------------------------------------------------------------------------- subtitles


def test_subtitles(S):
    (vtt,) = run_option("subtitle.to.vtt", [S["srt"]])
    v = vtt.read_text(encoding="utf-8")
    assert v.startswith("WEBVTT") and "00:00:01.000 --> 00:00:02.500" in v
    assert "01:00:05.250 --> 01:00:06.000" in v and "<i>Hello there.</i>" in v
    (srt,) = run_option("subtitle.to.srt", [S["vtt"]])
    s = srt.read_text(encoding="utf-8")
    assert s.startswith("1\n00:00:01,000 --> 00:00:02,000\nHi <b>you</b>")
    assert "2\n00:00:02,000 --> 00:00:03,500\nSecond cue" in s and "NOTE" not in s
    (txt,) = run_option("subtitle.to.txt", [S["srt"]])
    assert txt.read_text(encoding="utf-8").splitlines() == [
        "Hello there.", "General Kenobi!", "You are a bold one.", "Bye"]
    (txt2,) = run_option("subtitle.to.txt", [S["vtt"]])
    assert txt2.read_text(encoding="utf-8").splitlines() == ["Hi you", "Second cue"]


# --------------------------------------------------------------------------- archives


def test_extract_zip_blocks_zip_slip(S, tmp_path):
    src = tmp_path / "work" / "bundle.zip"
    src.parent.mkdir()
    shutil.copy(S["zip"], src)
    ctx_out = run_option("archive.extract", [src])
    (out,) = ctx_out
    assert out == src.parent / "bundle" and (out / "docs" / "readme.txt").read_text() == "inside zip"
    assert (out / "top.txt").exists()
    assert not (tmp_path / "evil.txt").exists() and not (src.parent / "evil.txt").exists()
    assert not list(out.rglob("evil*.txt"))


def test_extract_tar_blocks_traversal(S, tmp_path):
    src = tmp_path / "w" / "bundle.tar"
    src.parent.mkdir()
    shutil.copy(S["tar"], src)
    (out,) = run_option("archive.extract", [src])
    # single top-level folder gets lifted out instead of nesting bundle/folder
    assert out.name == "folder" and (out / "a.txt").read_bytes() == b"tar a"
    assert not (tmp_path / "evil.txt").exists() and not (src.parent / "evil.txt").exists()


def test_extract_single_file_tgz(S, tmp_path):
    src = tmp_path / "single.tar.gz"
    shutil.copy(S["tgz"], src)
    (out,) = run_option("archive.extract", [src])
    assert out.name == "only.txt" and out.read_bytes() == b"only"


def test_extract_plain_gz(tmp_path):
    import gzip

    src = tmp_path / "log.txt.gz"
    with gzip.open(src, "wb") as f:
        f.write(b"line one\n")
    (out,) = run_option("archive.extract", [src])
    assert out.name == "log.txt" and out.read_bytes() == b"line one\n"


def test_archive_conversions(S, tmp_path):
    src = tmp_path / "bundle.zip"
    shutil.copy(S["zip"], src)
    (tar,) = run_option("archive.to.tar", [src])
    with tarfile.open(tar) as t:
        names = t.getnames()
    assert "docs/readme.txt" in names and not any(".." in n for n in names)
    (tgz,) = run_option("archive.to.tgz", [tar])
    assert tgz.name == "bundle.tar.gz"
    with tarfile.open(tgz, "r:gz") as t:
        assert "top.txt" in t.getnames()
    (z2,) = run_option("archive.to.zip", [tgz])
    assert z2.name.startswith("bundle") and z2.suffix == ".zip"
    with zipfile.ZipFile(z2) as z:
        assert z.read("docs/readme.txt") == b"inside zip"


def test_pack_any(S, tmp_path):
    (z,) = run_option("any.to.zip", [S["folder"]])
    assert z.name == "My Folder.zip"
    with zipfile.ZipFile(z) as zf:
        assert {"My Folder/one.txt", "My Folder/sub/two.txt"} <= set(zf.namelist())
    a, b = tmp_path / "a.txt", tmp_path / "b.txt"
    a.write_text("a")
    b.write_text("b")
    (z2,) = run_option("any.to.zip", [a, b])
    assert z2.name == "Archive.zip"
    with zipfile.ZipFile(z2) as zf:
        assert sorted(zf.namelist()) == ["a.txt", "b.txt"]
    # A single file with its own conversions doesn't get packing wedges; unknown types do.
    from satsuma import registry
    assert not registry.get("any.to.zip").applies([a])
    blob = tmp_path / "a.bin"
    blob.write_bytes(b"\0\1")
    (one,) = run_option("any.to.zip", [blob])
    assert one.name == "a.zip"
    (tgz,) = run_option("any.to.tgz", [S["folder"]])
    with tarfile.open(tgz) as t:
        assert "My Folder/sub/two.txt" in t.getnames()
    (tar,) = run_option("any.to.tar", [a, b])
    assert tar.name == "Archive.tar"


def test_extract_7z_with_windows_tar(S, tmp_path):
    from satsuma.converters.archives import tar_exe

    try:
        exe = tar_exe()
    except UserError:
        pytest.skip("tar.exe not available")
    src_dir = tmp_path / "src"
    src_dir.mkdir()
    (src_dir / "seven.txt").write_text("7z!")
    (src_dir / "other.txt").write_text("x")
    arc = tmp_path / "pack.7z"
    r = subprocess.run([exe, "--format", "7zip", "-cf", str(arc), "-C", str(src_dir), "seven.txt", "other.txt"],
                       capture_output=True)
    if r.returncode != 0 or not arc.exists():
        pytest.skip("tar.exe cannot write 7z here")
    (out,) = run_option("archive.extract", [arc])
    assert (out / "seven.txt").read_text() == "7z!"


# --------------------------------------------------------------------------- PDF tools


def test_parse_ranges():
    from satsuma.tools.pdf_tools import page_set, parse_ranges

    assert parse_ranges("1-3,4, 5-9", 10) == [[0, 1, 2], [3], [4, 5, 6, 7, 8]]
    assert parse_ranges("8-", 10) == [[7, 8, 9]]
    assert parse_ranges("-2; last", 10) == [[0, 1], [9]]
    assert page_set("", 3) == [0, 1, 2]
    assert page_set("3,1,1", 3) == [0, 2]
    for bad in ("0", "11", "a-b", "", "1-12"):
        with pytest.raises(UserError):
            parse_ranges(bad, 10)


def test_compress(S):
    ctx_params = {"preset": "strong"}
    (out,) = run_option("pdf.compress", [S["bigpdf"]], ctx_params)
    assert out.name == "photos-compressed.pdf"
    assert out.stat().st_size < S["bigpdf"].stat().st_size * 0.7
    assert pdf_pages(out) == 2


def test_split_modes(S):
    (folder,) = run_option("pdf.split", [S["pdf"]], {"mode": "every"})
    assert sorted(p.name for p in folder.iterdir()) == [f"report page 0{i}.pdf" for i in (1, 2, 3)]
    (f2,) = run_option("pdf.split", [S["pdf"]], {"mode": "n", "n": 2})
    assert sorted(pdf_pages(p) for p in f2.iterdir()) == [1, 2]
    (f3,) = run_option("pdf.split", [S["pdf"]], {"mode": "ranges", "ranges": "1, 2-3"})
    assert sorted(p.name for p in f3.iterdir()) == ["report page 01.pdf", "report pages 2-3.pdf"]
    (one,) = run_option("pdf.split", [S["pdf"]], {"mode": "ranges", "ranges": "2-3"})
    assert one.is_file() and pdf_pages(one) == 2
    with pytest.raises(UserError):
        run_option("pdf.split", [S["pdf"]], {"mode": "n", "n": 5})


def test_merge_with_image(S, tmp_path):
    img = tmp_path / "pic.png"
    img.write_bytes(_png_bytes())
    order = [str(img), str(S["pdf"])]
    (out,) = run_option("pdf.merge", [S["pdf"], img], {"order": order})
    assert out.name == "Merged.pdf" and out.parent == tmp_path
    assert pdf_pages(out) == 4
    with pymupdf.open(str(out)) as d:
        assert "Quarterly Report" in d[1].get_text()


def test_rotate_extract_delete(S):
    (r,) = run_option("pdf.rotate", [S["pdf"]], {"angle": 90, "pages": "2-3"})
    with pymupdf.open(str(r)) as d:
        assert [p.rotation for p in d] == [0, 90, 90]
    (r2,) = run_option("pdf.rotate", [S["pdf"]], {"angle": 270})
    with pymupdf.open(str(r2)) as d:
        assert [p.rotation for p in d] == [270, 270, 270]
    (e,) = run_option("pdf.extract", [S["pdf"]], {"pages": "3,1"})
    with pymupdf.open(str(e)) as d:
        assert d.page_count == 2 and "Final page" in d[0].get_text()
    (x,) = run_option("pdf.delete", [S["pdf"]], {"pages": "2"})
    with pymupdf.open(str(x)) as d:
        assert d.page_count == 2 and "Pictures" not in "".join(p.get_text() for p in d)
    with pytest.raises(UserError):
        run_option("pdf.delete", [S["pdf"]], {"pages": "1-3"})


def test_protect_and_unlock(S):
    (p,) = run_option("pdf.protect", [S["pdf"]], {"password": "s3cret", "allow_copy": False})
    with pymupdf.open(str(p)) as d:
        assert d.needs_pass
        assert d.authenticate("s3cret")
        assert "AES" in (d.metadata.get("encryption") or "")
    with pytest.raises(UserError, match="Wrong password"):
        run_option("pdf.unlock", [p], {"password": "nope"})
    (u,) = run_option("pdf.unlock", [p], {"password": "s3cret"})
    with pymupdf.open(str(u)) as d:
        assert not d.needs_pass and not d.is_encrypted and d.page_count == 3
        assert "Quarterly Report" in d[0].get_text()
    with pytest.raises(UserError, match="not password protected"):
        run_option("pdf.unlock", [u], {"password": ""})


def test_watermark_and_page_numbers(S):
    (w,) = run_option("pdf.watermark", [S["pdf"]], {"text": "DRAFT", "opacity": 30, "angle": 45, "size": 50})
    with pymupdf.open(str(w)) as d:
        assert all("DRAFT" in p.get_text() for p in d)
    (n,) = run_option("pdf.page_numbers", [S["pdf"]], {"position": "bottom-right",
                                                       "format": "Page {n} of {total}"})
    with pymupdf.open(str(n)) as d:
        assert "Page 3 of 3" in d[2].get_text()
        words = [w for w in d[0].get_text("words") if w[4] == "Page"]
        assert words and words[0][1] > d[0].rect.height - 60 and words[0][0] > d[0].rect.width / 2


def _nfkc(s: str) -> str:
    """Arabic text may come back as presentation forms, NBSP spaces and bidi line breaks:
    normalise to NFKC with single spaces."""
    import unicodedata

    return " ".join(unicodedata.normalize("NFKC", s).split())


def _fonts_of(page, needle: str) -> set[str]:
    """Fonts of the spans holding characters of `needle` (stamped text lives in a Form XObject)."""
    chars = {c for c in _nfkc(needle) if not c.isspace() and not c.isascii()}
    return {s["font"] for b in page.get_text("dict")["blocks"] for line in b.get("lines", [])
            for s in line["spans"] if chars & set(_nfkc(s["text"]))}


def _win_font(name: str) -> bool:
    from satsuma.converters.documents import fonts_dir

    d = fonts_dir()
    return d is not None and (d / name).is_file()


SCRIPT_CASES = [
    # watermark text, page-number format, expected words in it, Windows font file, font name
    ("سري", "صفحة {n} من {total}", ["صفحة", "من"], "arial.ttf", "Arial"),
    ("机密", "第 {n} 页，共 {total} 页", ["第", "页"], "msyh.ttc", "YaHei"),
    ("गोपनीय", "पेज {n} का {total}", ["पेज", "का"], "Nirmala.ttf", "Nirmala"),
]


@pytest.mark.parametrize("text,fmt,words,font_file,font_name", SCRIPT_CASES, ids=["ar", "zh", "hi"])
def test_watermark_and_page_numbers_any_script(S, tmp_path, text, fmt, words, font_file, font_name):
    src = tmp_path / "mixed.pdf"
    with pymupdf.open(str(S["pdf"])) as d:
        d[1].set_rotation(90)  # stamping must work on rotated pages too
        d.save(str(src))
    (w,) = run_option("pdf.watermark", [src], {"text": text, "opacity": 30, "angle": 45, "size": 50})
    (n,) = run_option("pdf.page_numbers", [w], {"position": "bottom-right", "format": fmt, "size": 12})
    with pymupdf.open(str(n)) as d:
        for page in d:
            got = _nfkc(page.get_text())
            assert _nfkc(text) in got, (page.number, got)
            assert all(_nfkc(x) in got for x in words), (page.number, got)
            assert str(page.number + 1) in got
            fonts = _fonts_of(page, text) | _fonts_of(page, "".join(words))
            assert fonts and not any("helvetica" in f.lower() for f in fonts), fonts
            if _win_font(font_file):
                assert any(font_name in f for f in fonts), fonts
        # the page number sits in the visual bottom-right corner, also on the rotated page
        for page in d:
            r = page.rect
            spans = [s for b in page.get_text("dict")["blocks"] for line in b.get("lines", [])
                     for s in line["spans"] if _nfkc(words[0]) in _nfkc(s["text"])]
            assert spans, page.number
            box = pymupdf.Rect(spans[0]["bbox"])
            if page.rotation == 0:
                assert box.y0 > r.height - 60 and box.x1 > r.width / 2
    assert n.stat().st_size < 1_500_000  # fonts are subset, not embedded whole


def test_md_to_pdf_fallback_any_script(S, tmp_path, no_office):
    md = tmp_path / "arabic.md"
    md.write_text("# تقرير سري\n\nهذا نص عربي للاختبار.\n\n中文 机密 and हिन्दी गोपनीय text.\n",
                  encoding="utf-8")
    (pdf,) = run_option("doc.to.pdf", [md])
    with pymupdf.open(str(pdf)) as d:
        page = d[0]
        got = _nfkc(page.get_text())
        for word in ("تقرير", "سري", "عربي", "للاختبار", "机密", "गोपनीय", "text"):
            assert word in got, (word, got)
        fonts = _fonts_of(page, "سري عربي") | _fonts_of(page, "机密") | _fonts_of(page, "गोपनीय")
        assert fonts and not any("helvetica" in f.lower() for f in fonts), fonts
        if _win_font("arial.ttf"):
            assert any("Arial" in f for f in _fonts_of(page, "عربي"))
    assert pdf.stat().st_size < 1_500_000


def test_script_runs():
    from satsuma.converters.documents import localize_html, script_html, script_runs

    assert script_runs("Page 3") == [(None, "Page 3")]
    assert [s for s, _ in script_runs("سري 3 机密 गोपनीय")] == ["arabic", None, "cjk", None, "deva"]
    assert script_runs("क्ष") == [("deva", "क्ष")]  # virama stays in the run
    assert script_html("a<b") == "a&lt;b"
    html, used = localize_html("<html><head><title>机密</title></head><body><p>x 机密</p></body></html>")
    assert used == {"cjk"} and "<title>机密</title>" in html and '<span class="sz-cjk">机密</span>' in html


# --------------------------------------------------------------------------- dialogs


def test_dialogs_construct(S, tmp_path):
    from PySide6.QtWidgets import QApplication

    from satsuma.tools import pdf_tools as t

    app = QApplication.instance() or QApplication([])
    paths = [S["pdf"]]
    expected = {
        t.CompressDialog: {"preset"}, t.SplitDialog: {"mode", "n", "ranges"},
        t.RotateDialog: {"angle", "pages"}, t.ProtectDialog: {"password", "owner_password", "allow_print"},
        t.UnlockDialog: {"password"}, t.WatermarkDialog: {"text", "opacity", "angle", "size", "color"},
        t.PageNumbersDialog: {"position", "format", "start", "size"},
    }
    for cls, keys in expected.items():
        dlg = cls(paths)
        assert keys <= set(dlg.values()), cls
        dlg.deleteLater()
    pd = t.PagesDialog(paths, "Extract pages", "Extract", "e.g. 1-3")
    assert pd.values() == {"pages": ""}
    md = t.MergeDialog([S["pdf"], S["bigpdf"]])
    md.list.setCurrentRow(1)
    md._move(-1)
    assert md.values()["order"] == [str(S["bigpdf"]), str(S["pdf"])]
    assert t.CompressDialog(paths).values() == {"preset": "medium"}
    assert t.unlock_dialog([S["pdf"]], FakeCtx()) == {"password": ""}  # not locked: no prompt
    app.processEvents()


FAKE_AR = {
    "Watermark": "علامة مائية", "Add watermark": "إضافة علامة مائية", "Gray": "رمادي", "Red": "أحمر",
    "Colour": "اللون", "CONFIDENTIAL": "سري", "Page numbers": "أرقام الصفحات", "Top left": "أعلى اليسار",
    "Page {n} of {total}": "صفحة {n} من {total}", "Page {n}": "صفحة {n}", "Compress PDF": "ضغط PDF",
    "Medium — images at {dpi} DPI": "متوسط — الصور بدقة {dpi} نقطة/بوصة", "Every page": "كل صفحة",
    "Allow printing": "السماح بالطباعة", "{name} · {n} pages": "{name} · {n} صفحات",
    "90° clockwise": "90° باتجاه عقارب الساعة", "{name}-numbered": "{name}-مرقم",
}


def test_dialogs_in_arabic(S, monkeypatch):
    """Dialogs build under a right-to-left language: labels translated, values stay internal."""
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QApplication, QLabel

    from satsuma import i18n
    from satsuma.tools import pdf_tools as t

    app = QApplication.instance() or QApplication([])
    monkeypatch.setattr(i18n, "_catalog", i18n._catalog)  # restored after the test
    monkeypatch.setattr(i18n, "_current", i18n._current)
    i18n.set_language("ar")
    monkeypatch.setattr(i18n, "_catalog", dict(FAKE_AR))
    app.setLayoutDirection(Qt.RightToLeft)
    paths = [S["pdf"]]
    try:
        assert i18n.is_rtl() and registry.get("pdf.watermark").title == "علامة مائية"
        wd = t.WatermarkDialog(paths)
        assert "علامة مائية" in wd.windowTitle() and wd.primary_btn.text() == "إضافة علامة مائية"
        assert wd.text.text() == "سري" and wd.color.currentText() == "رمادي"
        assert "report.pdf · 3 صفحات" in [lbl.text() for lbl in wd.findChildren(QLabel)]  # subtitle
        wd.color.setCurrentIndex(wd.color.findData("Red"))
        assert wd.values()["color"] == "Red" and wd.color.currentText() == "أحمر"

        pd = t.PageNumbersDialog(paths)
        items = {pd.fmt.itemText(i): pd.fmt.itemData(i) for i in range(pd.fmt.count())}
        assert items["1"] == "{n}" and items["صفحة 1 من N"] == "صفحة {n} من {total}"
        assert items["Page 1 of N"] == "Page {n} of {total}"  # English stays available
        pd.fmt.setCurrentIndex(pd.fmt.findData("صفحة {n} من {total}"))
        pd.position.setCurrentIndex(pd.position.findData("top-left"))
        assert pd.position.currentText() == "أعلى اليسار"
        params = pd.values()
        assert params["position"] == "top-left" and params["format"] == "صفحة {n} من {total}"

        cd = t.CompressDialog(paths)
        assert cd.preset.currentText() == "متوسط — الصور بدقة 100 نقطة/بوصة" and cd.values() == {"preset": "medium"}
        assert t.SplitDialog(paths).values()["mode"] == "every"
        assert t.RotateDialog(paths).values()["angle"] == 90
        pr = t.ProtectDialog(paths)
        assert "السماح بالطباعة" in [cb.text() for cb in pr.checks.values()]
        for dlg in (wd, pd, cd, pr):
            assert dlg.layoutDirection() == Qt.RightToLeft
            dlg.adjustSize()
            dlg.deleteLater()

        # the Arabic format is what gets stamped; the output name is translated too
        (out,) = run_option("pdf.page_numbers", paths, params)
        assert out.name == "report-مرقم.pdf"
        with pymupdf.open(str(out)) as d:
            assert "صفحة 2 من 3" in _nfkc(d[1].get_text())
    finally:
        i18n.set_language("en")
        app.setLayoutDirection(Qt.LeftToRight)
    app.processEvents()
    assert registry.get("pdf.watermark").title == "Watermark"
